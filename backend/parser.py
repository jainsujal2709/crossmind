import io, os, csv
def parse(name, data):
    """Return [{'text','source'}] segments keeping page/slide/sheet/section info."""
    ext = os.path.splitext(name)[1].lower(); out = []
    if ext == ".pdf":
        import fitz
        for i, p in enumerate(fitz.open(stream=data, filetype="pdf"), 1):
            out.append({"text": p.get_text(), "source": f"{name} — page {i}"})
    elif ext == ".pptx":
        from pptx import Presentation
        for i, s in enumerate(Presentation(io.BytesIO(data)).slides, 1):
            out.append({"text": "\n".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame), "source": f"{name} — slide {i}"})
    elif ext == ".docx":
        from docx import Document
        sec, buf = "Introduction", []
        for p in Document(io.BytesIO(data)).paragraphs:
            if p.style.name.startswith("Heading"):
                out.append({"text": "\n".join(buf), "source": f"{name} — {sec}"}); sec, buf = p.text, []
            else: buf.append(p.text)
        out.append({"text": "\n".join(buf), "source": f"{name} — {sec}"})
    elif ext == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8", "ignore"))))
        out.append({"text": "\n".join(" ".join(r) for r in rows), "source": f"{name} — rows"})
    elif ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb: out.append({"text": "\n".join(" ".join(str(c) for c in r if c is not None) for r in ws.iter_rows(values_only=True)), "source": f"{name} — sheet {ws.title}"})
    elif ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=data)
        for ws in wb.sheets(): out.append({"text": "\n".join(" ".join(str(c) for c in ws.row_values(i) if c != "") for i in range(ws.nrows)), "source": f"{name} — sheet {ws.name}"})
    elif ext == ".txt":
        out.append({"text": data.decode("utf-8", "ignore"), "source": name})
    else:
        raise ValueError("Sorry, this file type isn't supported.")
    return [s for s in out if s["text"].strip()]
